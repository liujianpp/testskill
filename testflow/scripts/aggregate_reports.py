#!/usr/bin/env python3
"""多源执行结果汇总 → 统一 Markdown 报告。

源（缺哪个跳过哪个）:
    1. 25列 testcases.xlsx   —— 人工执行结果（实际结果/是否通过/执行时间/关联需求ID）
    2. allure-results 目录    —— 接口自动化（*-result.json 的 status）
    3. case-runner MD 报告    —— UI 执行（RESULT: PASS|FAIL|BLOCKED 行）
    4. miniprogram 报告       —— 小程序支线（run-with-reports.sh 生成的台账或 Session 批量报告）

用法:
    python aggregate_reports.py --excel artifacts/xx_testcases.xlsx \
        [--allure-dir allure-results] [--ui-report reports/case-run-x.md] \
        [--mp-report reports/miniprogram-executions.md] \
        [--output artifacts/unified-report.md]
"""
import argparse
import glob
import json
import re
from pathlib import Path

from openpyxl import load_workbook

PASS_TOKENS = {"是", "通过", "y", "yes", "pass", "passed", "true", "1"}
FAIL_TOKENS = {"否", "不通过", "未通过", "失败", "n", "no", "fail", "failed", "false", "0"}
UI_RESULT_RE = re.compile(r"(?:RESULT|结果)\**[:：]\s*\**\s*(PASS|FAIL|BLOCKED)", re.IGNORECASE)
# miniprogram 台账行：- **TC-MP-001** | PASS | exit=0 | spec=`...` | artifacts=`...`
MP_LEDGER_RE = re.compile(r"^-\s+\*\*(.+?)\*\*\s*\|\s*(PASS|FAIL|NOT_RUN)\s*\|", re.MULTILINE)
# miniprogram 批量报告「执行顺序」行：1. `tests/auth/auth.spec.js` — PASS（全角破折号）
MP_ORDER_RE = re.compile(r"^\d+\.\s+`(.+?)`\s+—\s+(PASS|FAIL|NOT_RUN)\s*$", re.MULTILINE)
MP_SESSION_RE = re.compile(r"^##\s+Session\b.*$", re.MULTILINE)


def cell_str(v) -> str:
    return "" if v is None else str(v).strip()


def read_excel_rows(path: str) -> list[dict]:
    """读取 25 列 Excel 的人工执行结果行。"""
    wb = load_workbook(path, data_only=True)
    ws = wb["测试用例"] if "测试用例" in wb.sheetnames else wb[wb.sheetnames[0]]
    headers = [cell_str(c.value) for c in ws[1]]

    def idx(name):
        return headers.index(name) + 1 if name in headers else None

    i_title, i_module = idx("用例标题"), idx("模块")
    i_actual, i_pass, i_req = idx("实际结果"), idx("是否通过"), idx("关联需求ID")

    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(cell_str(v) == "" for v in row):
            continue

        def g(i, _row=row):
            return cell_str(_row[i - 1]) if i else ""

        verdict = g(i_pass).lower()
        is_passed = True if verdict in PASS_TOKENS else (False if verdict in FAIL_TOKENS else None)
        rows.append({
            "title": g(i_title), "module": g(i_module), "req_id": g(i_req),
            "actual": g(i_actual), "is_passed": is_passed, "executed": is_passed is not None,
        })
    return rows


def read_allure_results(dirpath: str) -> list[dict]:
    """读取 allure-results 下的 *-result.json。"""
    rows = []
    for f in sorted(glob.glob(str(Path(dirpath) / "*-result.json"))):
        try:
            d = json.loads(Path(f).read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            continue
        rows.append({
            "name": str(d.get("name", "")),
            "status": str(d.get("status", "unknown")),
            "message": str((d.get("statusDetails") or {}).get("message", "")),
        })
    return rows


def read_ui_results(path: str) -> list[dict]:
    """解析 case-runner MD 报告中的 RESULT 行。"""
    rows = []
    for line in Path(path).read_text(encoding="utf-8", errors="replace").splitlines():
        m = UI_RESULT_RE.search(line)
        if m:
            rows.append({"result": m.group(1).upper(), "line": line.strip()})
    return rows


def read_mp_results(path: str) -> list[dict]:
    """解析 miniprogram-run 的报告（自动识别两种格式）。

    - 台账 miniprogram-executions.md：按 ``## Session`` 分段累积，取**最后一个** Session 段，
      解析 ``- **<用例ID>** | <STATUS> |`` 行（label 为用例 ID，可与 25 列编号对齐）。
    - Session 批量报告 miniprogram-report-<sessionId>.md：解析「执行顺序」的
      ``N. `<spec路径>` — <STATUS>`` 行（含 NOT_RUN，label 为 spec 路径）。
    """
    text = Path(path).read_text(encoding="utf-8", errors="replace")

    sessions = list(MP_SESSION_RE.finditer(text))
    if sessions:
        last = sessions[-1]
        chunk = text[last.end():]
        # 后续若再出现 Session 头则忽略（取最后一段）
        nxt = MP_SESSION_RE.search(chunk)
        if nxt:
            chunk = chunk[:nxt.start()]
        return [
            {"result": m.group(2), "label": m.group(1).strip()}
            for m in MP_LEDGER_RE.finditer(chunk)
        ]

    return [
        {"result": m.group(2), "label": m.group(1).strip()}
        for m in MP_ORDER_RE.finditer(text)
    ]


def _rate(p: int, f: int) -> str:
    """通过率 = 通过/(通过+失败)，分母为 0 时显示 '-'。"""
    return f"{p / (p + f) * 100:.1f}%" if (p + f) else "-"


def build_report(excel_rows, allure_rows, ui_rows, mp_rows=None) -> str:
    mp_rows = mp_rows or []

    # ── 统计 ──
    manual_pass = sum(1 for r in excel_rows if r["is_passed"] is True)
    manual_fail = sum(1 for r in excel_rows if r["is_passed"] is False)
    manual_pending = sum(1 for r in excel_rows if not r["executed"])

    allure_pass = sum(1 for r in allure_rows if r["status"] == "passed")
    allure_fail = sum(1 for r in allure_rows if r["status"] in ("failed", "broken"))

    ui_pass = sum(1 for r in ui_rows if r["result"] == "PASS")
    ui_fail = sum(1 for r in ui_rows if r["result"] == "FAIL")
    ui_blocked = sum(1 for r in ui_rows if r["result"] == "BLOCKED")

    mp_pass = sum(1 for r in mp_rows if r["result"] == "PASS")
    mp_fail = sum(1 for r in mp_rows if r["result"] == "FAIL")
    mp_not_run = sum(1 for r in mp_rows if r["result"] == "NOT_RUN")

    lines = ["# 测试统一报告", "", "## 1. 执行总览", ""]
    lines.append("| 来源 | 用例数 | 通过 | 失败 | 阻塞 | 待执行 | 通过率 |")
    lines.append("|---|---|---|---|---|---|---|")
    lines.append(f"| 人工用例(25列Excel) | {len(excel_rows)} | {manual_pass} | {manual_fail} | - | {manual_pending} | {_rate(manual_pass, manual_fail)} |")
    lines.append(f"| 接口自动化(Allure) | {len(allure_rows)} | {allure_pass} | {allure_fail} | - | - | {_rate(allure_pass, allure_fail)} |")
    lines.append(f"| UI执行(case-runner) | {len(ui_rows)} | {ui_pass} | {ui_fail} | {ui_blocked} | - | {_rate(ui_pass, ui_fail)} |")
    lines.append(f"| 小程序执行(miniprogram) | {len(mp_rows)} | {mp_pass} | {mp_fail} | - | {mp_not_run} | {_rate(mp_pass, mp_fail)} |")
    lines.append("")
    lines.append("> 通过率 = 通过/(通过+失败)；阻塞/未执行单列不计入分母。")
    lines.append("")

    # ── 需求覆盖 ──
    lines += ["## 2. 需求覆盖对照", ""]
    reqs = {}
    for r in excel_rows:
        key = r["req_id"] or "（未关联）"
        reqs.setdefault(key, []).append(r)
    if reqs:
        lines.append("| 需求ID | 用例数 | 通过 | 失败 | 待执行 | 通过率 |")
        lines.append("|---|---|---|---|---|---|")
        for key in sorted(reqs):
            group = reqs[key]
            p = sum(1 for r in group if r["is_passed"] is True)
            f = sum(1 for r in group if r["is_passed"] is False)
            pending = sum(1 for r in group if not r["executed"])
            lines.append(f"| {key} | {len(group)} | {p} | {f} | {pending} | {_rate(p, f)} |")
    else:
        lines.append("（无 25 列 Excel 数据，需求覆盖待执行）")
    lines.append("")

    # ── 明细 ──
    lines += ["## 3. 执行明细", ""]
    lines += ["### 3.1 人工用例（25列 Excel）", ""]
    if not excel_rows:
        lines.append("（未提供 25 列 Excel）")
    else:
        lines.append("| 用例标题 | 模块 | 需求ID | 结果 | 实际结果 |")
        lines.append("|---|---|---|---|---|")
        for r in excel_rows:
            verdict = "待执行" if not r["executed"] else ("通过" if r["is_passed"] else "失败")
            lines.append(f"| {r['title']} | {r['module']} | {r['req_id'] or '-'} | {verdict} | {r['actual'] or '-'} |")
    lines.append("")

    lines += ["### 3.2 接口自动化（Allure）", ""]
    if not allure_rows:
        lines.append("（未提供 allure-results 目录或无结果文件）")
    else:
        failed = [r for r in allure_rows if r["status"] in ("failed", "broken")]
        lines.append(f"共 {len(allure_rows)} 条，失败/中断 {len(failed)} 条。失败明细：")
        lines.append("")
        for r in failed:
            lines.append(f"- **{r['name']}** [{r['status']}]：{r['message'] or '（无失败信息）'}")
    lines.append("")

    lines += ["### 3.3 UI 执行（case-runner）", ""]
    if not ui_rows:
        lines.append("（UI 报告未生成）")
    else:
        for r in ui_rows:
            lines.append(f"- {r['result']}：{r['line']}")
    lines.append("")

    lines += ["### 3.4 小程序执行（miniprogram）", ""]
    if not mp_rows:
        lines.append("（小程序报告未生成）")
    else:
        for r in mp_rows:
            lines.append(f"- {r['result']}：{r['label']}")
    lines.append("")

    # ── 风险与建议 ──
    lines += ["## 4. 风险与建议", ""]
    risks = []
    if manual_fail:
        mods = {r["module"] for r in excel_rows if r["is_passed"] is False and r["module"]}
        risks.append(f"- 人工用例失败 {manual_fail} 条，集中在模块：{', '.join(sorted(mods)) or '-'}")
    if allure_fail:
        names = [r["name"] for r in allure_rows if r["status"] in ("failed", "broken")]
        risks.append(f"- 接口自动化失败 {allure_fail} 条：{', '.join(names[:10])}{'…' if len(names) > 10 else ''}")
    if ui_fail or ui_blocked:
        risks.append(f"- UI 执行失败 {ui_fail} 条、阻塞 {ui_blocked} 条（详见截图与报告）")
    if mp_fail:
        labels = [r["label"] for r in mp_rows if r["result"] == "FAIL"]
        risks.append(f"- 小程序执行失败 {mp_fail} 条：{', '.join(labels[:10])}{'…' if len(labels) > 10 else ''}（详见 trace viewer 与 artifacts）")
    if mp_not_run:
        risks.append(f"- 小程序有 {mp_not_run} 条未执行（STOP_ON_FAIL 停跑或 --from 续跑场景，可用 suite --from 续跑）")
    weak = [k for k, g in reqs.items() if k != "（未关联）" and all(not r["executed"] or r["is_passed"] is False for r in g)]
    if weak:
        risks.append(f"- 覆盖薄弱需求（无用例通过）：{', '.join(sorted(weak))}")
    if manual_pending:
        risks.append(f"- 尚有 {manual_pending} 条人工用例待执行")
    lines.extend(risks if risks else ["- 暂无显著风险"])

    return "\n".join(lines) + "\n"


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="多源执行结果汇总 → 统一 Markdown 报告")
    ap.add_argument("--excel", default=None, help="25列 testcases.xlsx 路径（默认: cwd 的 artifacts/ 下首个 *_testcases.xlsx）")
    ap.add_argument("--allure-dir", default=None, help="allure-results 目录（缺省则跳过该节）")
    ap.add_argument("--ui-report", default=None, help="case-runner MD 报告路径（缺省则跳过该节）")
    ap.add_argument("--mp-report", default=None,
                    help="miniprogram 报告路径（台账 miniprogram-executions.md 或 Session 批量报告；缺省则跳过该节）")
    ap.add_argument("--output", default=None, help="输出路径（默认: Excel 同目录 unified-report.md）")
    args = ap.parse_args(argv)

    # --excel 未提供时，从 cwd 的 artifacts/ 子目录（含 cwd 本身）推导
    if not args.excel:
        candidates = sorted(
            list(Path.cwd().glob("*_testcases.xlsx"))
            + list((Path.cwd() / "artifacts").glob("*_testcases.xlsx"))
        )
        if not candidates:
            ap.error("--excel 未提供且 artifacts/ 下无 *_testcases.xlsx")
        args.excel = str(candidates[0])
        print(f"使用推导 Excel: {args.excel}")

    excel_rows = read_excel_rows(args.excel)
    allure_rows = read_allure_results(args.allure_dir) if args.allure_dir else []
    ui_rows = read_ui_results(args.ui_report) if args.ui_report and Path(args.ui_report).exists() else []
    mp_rows = read_mp_results(args.mp_report) if args.mp_report and Path(args.mp_report).exists() else []

    report = build_report(excel_rows, allure_rows, ui_rows, mp_rows)

    # 注意：output 默认依赖 args.excel 的 parent，必须在 excel 推导之后计算
    output = args.output or str(Path(args.excel).parent / "unified-report.md")
    Path(output).write_text(report, encoding="utf-8")
    print(f"统一报告已生成: {output}")
    print(f"  人工用例: {len(excel_rows)} 行 | Allure: {len(allure_rows)} 条 | UI: {len(ui_rows)} 条 | 小程序: {len(mp_rows)} 条")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
