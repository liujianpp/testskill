import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

SCRIPTS_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SCRIPTS_DIR))

from convert_25col_to_ui_xlsx import RUNNER_HEADERS, convert

HEADERS_25 = [
    "系统/项目", "模块", "功能", "测试项/测试点", "用例标题",
    "前置条件", "涉及相关数据库表", "测试数据准备", "测试步骤", "预期结果",
    "验证SQL", "实际结果", "是否通过", "预估工时", "执行时间",
    "测试进度", "用例类型", "用例优先级", "是否回归用例", "是否阻塞上线",
    "关联需求ID", "API URL", "BUG数", "BUG未解决数", "用例版本",
]


def build_25col(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"
    ws.append(HEADERS_25)
    for r in rows:
        ws.append(r)
    wb.save(path)


def row_for(**kw):
    """构造一行 25 列数据，未给字段的列留空。"""
    base = {h: "" for h in HEADERS_25}
    base.update(kw)
    return [base[h] for h in HEADERS_25]


def load_out(path):
    wb = load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    return rows[0], rows[1:]


def test_standard_mapping(tmp_path):
    src = tmp_path / "in.xlsx"
    out = tmp_path / "out.xlsx"
    build_25col(src, [
        row_for(**{
            "模块": "登录", "功能": "账号登录", "测试项/测试点": "TC-001",
            "用例标题": "正确账号密码登录", "用例优先级": "P0",
            "前置条件": "账号已注册", "测试步骤": "1.打开登录页\n2.输入账号密码\n3.点击登录",
            "测试数据准备": "user=admin", "预期结果": "登录成功跳转首页",
            "用例类型": "正向",
        }),
    ])
    stats = convert(str(src), str(out))
    headers, rows = load_out(out)
    assert list(headers) == RUNNER_HEADERS == [
        "用例编号", "用例名称", "功能模块", "优先级", "前置条件",
        "测试步骤", "测试数据", "预期结果",
    ]
    assert len(rows) == 1
    r = rows[0]
    assert r[0] == "TC-001"           # 用例编号 ← 测试项/测试点
    assert r[1] == "正确账号密码登录"   # 用例名称 ← 用例标题
    assert r[2] == "登录"              # 功能模块 ← 模块
    assert r[3] == "P0"               # 优先级 ← 用例优先级
    assert r[4] == "账号已注册"
    assert r[5] == "1.打开登录页\n2.输入账号密码\n3.点击登录"  # 换行保留
    assert r[6] == "user=admin"       # 测试数据 ← 测试数据准备
    assert r[7] == "登录成功跳转首页"
    assert stats == {"total": 1, "exported": 1, "skipped_api": 0,
                     "skipped_type": 0, "skipped_dup": 0, "skipped_empty": 0}


def test_api_url_rows_excluded(tmp_path):
    src, out = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    build_25col(src, [
        row_for(**{"测试项/测试点": "TC-A", "用例标题": "UI用例"}),
        row_for(**{"测试项/测试点": "TC-B", "用例标题": "API用例", "API URL": "/api/login"}),
    ])
    stats = convert(str(src), str(out))
    _, rows = load_out(out)
    assert [r[0] for r in rows] == ["TC-A"]
    assert stats["skipped_api"] == 1
    assert stats["exported"] == 1


def test_fallback_id_generation(tmp_path):
    src, out = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    build_25col(src, [
        row_for(**{"模块": "登录", "功能": "账号登录", "用例标题": "无编号用例"}),
    ])
    convert(str(src), str(out))
    _, rows = load_out(out)
    assert rows[0][0] == "登录-账号登录-1"   # 模块-功能-行号


def test_dedup_keeps_first(tmp_path):
    src, out = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    build_25col(src, [
        row_for(**{"测试项/测试点": "TC-DUP", "用例标题": "第一条"}),
        row_for(**{"测试项/测试点": "TC-DUP", "用例标题": "第二条"}),
    ])
    stats = convert(str(src), str(out))
    _, rows = load_out(out)
    assert len(rows) == 1
    assert rows[0][1] == "第一条"
    assert stats["skipped_dup"] == 1


def test_types_filter_on_coverage_dim(tmp_path):
    src, out = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    build_25col(src, [
        row_for(**{"测试项/测试点": "TC-1", "用例标题": "正向", "用例类型": "正向"}),
        row_for(**{"测试项/测试点": "TC-2", "用例标题": "安全", "用例类型": "安全"}),
    ])
    stats = convert(str(src), str(out), include_types={"正向"})
    _, rows = load_out(out)
    assert [r[0] for r in rows] == ["TC-1"]
    assert stats["skipped_type"] == 1


def test_missing_source_column_no_crash(tmp_path):
    """缺「测试数据准备」列时该输出列留空，不报错。"""
    src, out = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"
    heads = [h for h in HEADERS_25 if h != "测试数据准备"]
    ws.append(heads)
    base = {h: "" for h in heads}
    base.update({"测试项/测试点": "TC-X", "用例标题": "缺列用例"})
    ws.append([base[h] for h in heads])
    wb.save(src)
    convert(str(src), str(out))
    _, rows = load_out(out)
    assert rows[0][0] == "TC-X"
    assert rows[0][6] in ("", None)  # openpyxl 不持久化空串，读回为 None


def test_empty_rows_skipped(tmp_path):
    src, out = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"
    ws.append(HEADERS_25)
    ws.append([""] * 25)
    ws.append(row_for(**{"测试项/测试点": "TC-E", "用例标题": "有效行"}))
    wb.save(src)
    stats = convert(str(src), str(out))
    assert stats["skipped_empty"] == 1
    assert stats["exported"] == 1


def test_wrong_sheet_errors(tmp_path):
    src, out = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    wb = Workbook()
    wb.active.title = "别的"
    wb.save(src)
    with pytest.raises(SystemExit, match="测试用例"):
        convert(str(src), str(out))
