#!/usr/bin/env python3
"""Validate seed-test-data outputs: Postman, SQL, Excel, and Python seed scripts."""

from __future__ import annotations

import ast
import json
import re
import sys
from pathlib import Path

POSTMAN_SCHEMA = "https://schema.getpostman.com/json/collection/v2.1.0/collection.json"
REQUIRED_FOLDERS = ("造数", "回滚")
SUMMARY_SHEET = "造数清单"
PYTHON_SUFFIXES = ("-api-seed.py", "-db-seed.py")


def fail(msg: str) -> None:
    print(f"FAIL: {msg}")
    sys.exit(1)


def ok(msg: str) -> None:
    print(f"OK: {msg}")


def validate_collection(path: Path) -> None:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        fail(f"{path.name}: invalid JSON ({e})")

    schema = (data.get("info") or {}).get("schema") or ""
    if POSTMAN_SCHEMA not in schema:
        fail(f"{path.name}: info.schema must be Collection v2.1")

    names = [item.get("name") for item in data.get("item") or [] if isinstance(item, dict)]
    for folder in REQUIRED_FOLDERS:
        if folder not in names:
            fail(f"{path.name}: missing folder '{folder}'")

    seed_folder = next(i for i in data["item"] if i.get("name") == "造数")
    if not seed_folder.get("item"):
        fail(f"{path.name}: folder '造数' has no requests")

    ok(path.name)


def validate_environment(path: Path) -> None:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        fail(f"{path.name}: invalid JSON ({e})")
    keys = {v.get("key") for v in data.get("values") or [] if isinstance(v, dict)}
    if "baseUrl" not in keys:
        fail(f"{path.name}: environment missing baseUrl")
    ok(path.name)


def _sql_deletes_unscoped(sql: str) -> bool:
    for line in sql.upper().splitlines():
        stripped = line.strip()
        if stripped.startswith("--") or not stripped:
            continue
        if "DELETE FROM" in stripped and "WHERE" not in stripped:
            return True
        if stripped.startswith("TRUNCATE") or stripped.startswith("DROP "):
            return True
    return False


def validate_sql_pair(seed: Path, rollback: Path) -> None:
    seed_sql = seed.read_text(encoding="utf-8")
    rb_sql = rollback.read_text(encoding="utf-8")
    if "INSERT" not in seed_sql.upper():
        fail(f"{seed.name}: seed SQL has no INSERT")
    if "DELETE" not in rb_sql.upper():
        fail(f"{rollback.name}: rollback SQL has no DELETE")
    if _sql_deletes_unscoped(rb_sql):
        fail(f"{rollback.name}: rollback has TRUNCATE/DROP or DELETE without WHERE")
    ok(f"{seed.name} + {rollback.name}")


def validate_excel(path: Path) -> None:
    try:
        from openpyxl import load_workbook
    except ImportError:
        fail("缺少 openpyxl，请先执行: pip install openpyxl")

    try:
        wb = load_workbook(path, read_only=True)
    except Exception as e:  # noqa: BLE001
        fail(f"{path.name}: cannot open xlsx ({e})")

    if SUMMARY_SHEET not in wb.sheetnames:
        wb.close()
        fail(f"{path.name}: missing sheet '{SUMMARY_SHEET}'")
    wb.close()
    ok(path.name)


def _python_has_dangerous_sql(source: str) -> str | None:
    """Return error reason if source embeds dangerous SQL literals."""
    if re.search(r"\bTRUNCATE\b", source, re.IGNORECASE):
        return "contains TRUNCATE"
    if re.search(r"\bDROP\s+(TABLE|DATABASE|SCHEMA)\b", source, re.IGNORECASE):
        return "contains DROP TABLE/DATABASE"

    for match in re.finditer(
        r"""(?:["'])([^"']*DELETE\s+FROM[^"']*)(?:["'])""",
        source,
        re.IGNORECASE | re.DOTALL,
    ):
        fragment = match.group(1)
        if "WHERE" not in fragment.upper():
            return "DELETE without WHERE in string literal"
    return None


def validate_python_seed(path: Path) -> None:
    source = path.read_text(encoding="utf-8")
    try:
        tree = ast.parse(source, filename=str(path))
    except SyntaxError as e:
        fail(f"{path.name}: invalid Python ({e})")

    func_names = {
        n.name
        for n in tree.body
        if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))
    }
    # Accept do_seed/do_rollback or any def whose name contains seed/rollback
    has_seed = "do_seed" in func_names or any("seed" in n for n in func_names)
    has_rollback = "do_rollback" in func_names or any("rollback" in n for n in func_names)
    if not has_seed:
        # Also accept CLI string dispatch
        if "seed" not in source:
            fail(f"{path.name}: missing seed action")
    if not has_rollback:
        if "rollback" not in source:
            fail(f"{path.name}: missing rollback action")

    if "seed" not in source or "rollback" not in source:
        fail(f"{path.name}: source must mention both seed and rollback")

    danger = _python_has_dangerous_sql(source)
    if danger:
        fail(f"{path.name}: {danger}")

    ok(path.name)


def collect_targets(args: list[str]) -> list[Path]:
    paths: list[Path] = []
    for raw in args:
        p = Path(raw)
        if p.is_dir():
            paths.extend(sorted(p.glob("*.json")))
            paths.extend(sorted(p.glob("*.sql")))
            paths.extend(sorted(p.glob("*.xlsx")))
            paths.extend(sorted(p.glob("*.py")))
            # Python seed scripts live under script/ (HTTP only; legacy *-db-seed.py may still exist)
            paths.extend(sorted(p.glob("script/*-api-seed.py")))
            paths.extend(sorted(p.glob("script/*-db-seed.py")))
        elif p.is_file():
            paths.append(p)
        else:
            fail(f"not found: {raw}")
    return paths


def _is_seed_python(path: Path) -> bool:
    name = path.name
    return any(name.endswith(suf) for suf in PYTHON_SUFFIXES)


def main() -> None:
    if len(sys.argv) < 2:
        fail("usage: validate_artifacts.py <file-or-dir>...")

    paths = collect_targets(sys.argv[1:])
    collections = [p for p in paths if p.name.endswith("postman_collection.json")]
    envs = [p for p in paths if p.name.endswith("postman_environment.json")]
    seeds = [p for p in paths if p.name.endswith("01_seed.sql")]
    rollbacks = [p for p in paths if p.name.endswith("02_rollback.sql")]
    excels = [p for p in paths if p.suffix.lower() == ".xlsx"]
    py_seeds = [p for p in paths if _is_seed_python(p)]

    if not collections and not seeds and not excels and not py_seeds:
        fail("no collection JSON, 01_seed.sql, .xlsx, or script/*-api-seed.py found")

    for c in collections:
        validate_collection(c)
    for e in envs:
        validate_environment(e)

    rb_by_stem = {}
    for rb in rollbacks:
        key = rb.name.replace("-02_rollback.sql", "").replace("02_rollback.sql", "")
        rb_by_stem[key] = rb

    for seed in seeds:
        key = seed.name.replace("-01_seed.sql", "").replace("01_seed.sql", "")
        rb = rb_by_stem.get(key)
        if rb is None:
            if len(seeds) == 1 and len(rollbacks) == 1:
                rb = rollbacks[0]
            else:
                fail(f"{seed.name}: missing matching 02_rollback.sql")
        validate_sql_pair(seed, rb)

    for x in excels:
        validate_excel(x)

    for py in py_seeds:
        validate_python_seed(py)

    print("OK")


if __name__ == "__main__":
    main()
