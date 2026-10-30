#!/usr/bin/env python3
"""Export seed-test-data JSON to Excel (.xlsx).

Requires: pip install openpyxl
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
except ImportError:
    print("缺少 openpyxl，请先执行: pip install openpyxl", file=sys.stderr)
    raise SystemExit(1)

SUMMARY_HEADERS = [
    "序号",
    "用途",
    "画像",
    "资源",
    "链路",
    "关键数据",
    "查找方式",
    "回滚方式",
]

SUMMARY_WIDTHS = {
    "A": 8,
    "B": 18,
    "C": 10,
    "D": 20,
    "E": 14,
    "F": 36,
    "G": 32,
    "H": 36,
}


def _safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = re.sub(r'[\\/*?:\[\]]', "_", name).strip() or "sheet"
    cleaned = cleaned[:31]
    base = cleaned
    n = 2
    while cleaned in used:
        suffix = f"_{n}"
        cleaned = (base[: 31 - len(suffix)] + suffix)[:31]
        n += 1
    used.add(cleaned)
    return cleaned


def load_data(path: Path) -> dict:
    with path.open(encoding="utf-8") as f:
        data = json.load(f)
    if not isinstance(data, dict):
        raise ValueError("JSON 根必须为对象")
    if "summary" not in data or not isinstance(data["summary"], list):
        raise ValueError("JSON 需包含非空 summary 数组")
    if not data["summary"]:
        raise ValueError("summary 不能为空")
    if "sheets" not in data or not isinstance(data["sheets"], list) or not data["sheets"]:
        raise ValueError("JSON 需包含非空 sheets 数组")
    return data


def _style_header_row(ws, col_count: int) -> None:
    header_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    for col in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.alignment = wrap


def _write_summary(ws, summary: list) -> None:
    wrap = Alignment(wrap_text=True, vertical="top")
    for col, label in enumerate(SUMMARY_HEADERS, start=1):
        ws.cell(row=1, column=col, value=label)
    _style_header_row(ws, len(SUMMARY_HEADERS))

    for row_idx, item in enumerate(summary, start=2):
        if not isinstance(item, dict):
            raise ValueError(f"summary[{row_idx - 2}] 必须为对象")
        values = [
            row_idx - 1,
            item.get("purpose", ""),
            item.get("profile", ""),
            item.get("resource", ""),
            item.get("chain", ""),
            item.get("key_data", ""),
            item.get("lookup", ""),
            item.get("rollback", ""),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value="" if value is None else str(value) if col > 1 else value)
            cell.alignment = wrap

    ws.freeze_panes = "A2"
    for letter, width in SUMMARY_WIDTHS.items():
        ws.column_dimensions[letter].width = width


def _write_detail(ws, sheet: dict) -> None:
    wrap = Alignment(wrap_text=True, vertical="top")
    headers = sheet.get("headers")
    rows = sheet.get("rows")
    if not isinstance(headers, list) or not headers:
        raise ValueError(f"sheet '{sheet.get('name')}' 需要非空 headers")
    if not isinstance(rows, list):
        raise ValueError(f"sheet '{sheet.get('name')}' 的 rows 必须为数组")

    full_headers = ["画像"] + [str(h) for h in headers]
    for col, label in enumerate(full_headers, start=1):
        ws.cell(row=1, column=col, value=label)
    _style_header_row(ws, len(full_headers))

    expected_len = 1 + len(headers)
    for row_idx, row in enumerate(rows, start=2):
        if not isinstance(row, list):
            raise ValueError(f"sheet '{sheet.get('name')}' rows[{row_idx - 2}] 必须为数组")
        if len(row) != expected_len:
            raise ValueError(
                f"sheet '{sheet.get('name')}' rows[{row_idx - 2}] 长度应为 {expected_len} "
                f"(画像 + {len(headers)} 字段)，实际 {len(row)}"
            )
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col, value="" if value is None else str(value))
            cell.alignment = wrap

    ws.freeze_panes = "A2"
    for i in range(1, len(full_headers) + 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = 14 if i == 1 else 20


def export_excel(data: dict, output: Path) -> None:
    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "造数清单"
    _write_summary(summary_ws, data["summary"])

    used_names = {"造数清单"}
    for sheet in data["sheets"]:
        if not isinstance(sheet, dict):
            raise ValueError("sheets 元素必须为对象")
        name = _safe_sheet_name(str(sheet.get("name") or "resource"), used_names)
        ws = wb.create_sheet(title=name)
        _write_detail(ws, sheet)

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description="将造数 JSON 导出为 Excel")
    parser.add_argument("input_json", type=Path, help="输入 JSON 路径")
    parser.add_argument("output_xlsx", type=Path, help="输出 .xlsx 路径")
    args = parser.parse_args()

    try:
        data = load_data(args.input_json)
        export_excel(data, args.output_xlsx)
    except Exception as exc:  # noqa: BLE001 — CLI surface
        print(f"导出失败: {exc}", file=sys.stderr)
        return 1

    print(f"已导出 Excel: {args.output_xlsx.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
