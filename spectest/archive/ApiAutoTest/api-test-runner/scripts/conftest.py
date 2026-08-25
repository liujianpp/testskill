"""
Pytest fixtures for CSV/Excel-driven API tests.
Provides auth_cookie (session-scoped) for login-dependent cases.
Config from config/dependencies.yaml.
run_tests.py chdirs to project root before pytest, so project_root = cwd().
"""

import csv
import logging
import os
from pathlib import Path

import pytest

# 请求/响应日志：运行时可看到每个请求的 method、url、headers、body 与响应状态、body
logging.basicConfig(level=logging.INFO, format="%(message)s")
logging.getLogger("urllib3").setLevel(logging.WARNING)
import requests
import yaml

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

# 用于结果写回：用例ID -> PASS/FAIL；pytest_sessionfinish 时写回 CSV 或 Excel
_csv_results: dict[str, str] = {}
_case_path: Path | None = None
_is_excel: bool = False


def _project_root() -> Path:
    """Project root: run_tests.py chdirs here before invoking pytest."""
    return Path.cwd()


def _project_root_from_file() -> Path:
    """从 conftest 文件位置推导项目根（scripts -> api-test-runner -> skills -> .cursor -> root）。"""
    return Path(__file__).resolve().parent.parent.parent.parent.parent


def _load_dotenv():
    """加载 config/.env 到 os.environ，供前置依赖等使用。"""
    root = _project_root_from_file()
    env_file = root / "config" / ".env"
    if env_file.exists():
        for line in env_file.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            if "=" in line:
                k, v = line.split("=", 1)
                k, v = k.strip(), v.strip()
                if k:
                    os.environ[k] = v


# 导入时即加载 .env，保证 session fixture 执行前已有环境变量
_load_dotenv()


def _load_dependencies() -> dict:
    """Load dependencies.yaml, resolve env vars."""
    config_path = _project_root() / "config" / "dependencies.yaml"
    if config_path.exists():
        raw = config_path.read_text(encoding="utf-8")
        for key, val in os.environ.items():
            raw = raw.replace(f"${{{key}}}", str(val))
        return yaml.safe_load(raw) or {}
    return {}


def _resolve_url(url_template: str) -> str:
    base = os.environ.get("BASE_URL", "http://localhost:8080")
    return url_template.replace("${BASE_URL}", base)


def _extract_json_path(obj: dict, path: str):
    """Simple JSONPath-like extract: $.data.token -> obj['data']['token']"""
    if not path or not path.startswith("$."):
        return None
    keys = path[2:].split(".")
    cur = obj
    for k in keys:
        cur = cur.get(k) if isinstance(cur, dict) else None
        if cur is None:
            break
    return cur


def _resolve_body(body: dict) -> dict:
    out = {}
    for k, v in (body or {}).items():
        if isinstance(v, str) and v.startswith("${") and v.endswith("}"):
            out[k] = os.environ.get(v[2:-1], "")
        else:
            out[k] = v
    return out


def _should_run(row: dict) -> bool:
    """是否运行：是/1/Y/true 或空视为执行，否/0/N/false 视为跳过。"""
    v = (row.get("是否运行") or "").strip().lower()
    if not v:
        return True
    if v in ("是", "1", "y", "true", "yes"):
        return True
    if v in ("否", "0", "n", "false", "no"):
        return False
    return True


def _load_excel_rows(case_path: Path) -> list[dict]:
    """Load test case rows from Excel (.xlsx). First row = headers. Filter by 接口路径 and 是否运行."""
    if load_workbook is None:
        raise ImportError("openpyxl is required for Excel. Install with: pip install openpyxl")
    wb = load_workbook(case_path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        wb.close()
        return []
    # 表头规范化：去空格、BOM
    keys = [str(c).lstrip("\ufeff").strip() if c is not None else "" for c in header]
    cases = []
    for row in rows_iter:
        case = {}
        for i, k in enumerate(keys):
            if not k:
                continue
            v = row[i] if i < len(row) else None
            case[k] = "" if v is None else str(v).strip()
        if not case.get("接口路径"):
            continue
        if not _should_run(case):
            continue
        cases.append(case)
    wb.close()
    return cases


def _load_case_rows(case_path: Path) -> list[dict]:
    """Load test case rows from CSV or Excel. Sets _case_path and _is_excel."""
    global _case_path, _is_excel
    _case_path = case_path
    _is_excel = case_path.suffix.lower() in (".xlsx", ".xls")
    if not case_path or not case_path.exists():
        return []
    if _is_excel:
        return _load_excel_rows(case_path)
    # CSV
    with open(case_path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames:
            return []
        cases = []
        for row in reader:
            case = {k.lstrip("\ufeff").strip(): (v or "") for k, v in row.items()}
            if not case.get("接口路径"):
                continue
            if not _should_run(case):
                continue
            cases.append(case)
    return cases


@pytest.fixture(scope="session")
def auth_cookie():
    """
    按 dependencies.yaml 执行所有前置依赖，返回 { 依赖名: { 提取字段: 值 } }。
    CSV 中「前置依赖」填依赖名（如 user_login），用例会拿到该依赖的产出并按「注入方式」注入。
    """
    deps_cfg = ( _load_dependencies().get("dependencies") ) or {}
    if not deps_cfg:
        return {}

    outputs = {}
    for dep_name, dep_config in deps_cfg.items():
        url = _resolve_url(dep_config.get("url", ""))
        if not url:
            continue
        body = _resolve_body(dep_config.get("body"))
        try:
            resp = requests.request(
                dep_config.get("method", "POST"),
                url,
                json=body,
                timeout=10,
            )
            resp.raise_for_status()
        except Exception:
            outputs[dep_name] = None
            continue

        data = resp.json() if "application/json" in (resp.headers.get("content-type") or "").lower() else {}
        extract = dep_config.get("extract") or {}
        one = {}
        for key, path in extract.items():
            if path and isinstance(path, str):
                one[key] = _extract_json_path(data, path)
        outputs[dep_name] = one if one else None

    return outputs


def _get_first_available_token(auth_cookie: dict) -> str | None:
    """从任意前置依赖产出中取 token，兼容 user_login 等。"""
    if not auth_cookie:
        return None
    for _name, out in auth_cookie.items():
        if out and isinstance(out, dict) and out.get("token"):
            return out["token"]
    return None


def pytest_addoption(parser):
    parser.addoption("--csv", action="store", default=None, help="CSV or Excel (.xlsx) test case file path")


def pytest_runtest_makereport(item, call):
    """Collect pass/fail per 用例ID for writing back to CSV/Excel."""
    if call.when != "call":
        return
    if "row" not in getattr(item, "funcargs", {}):
        return
    row = item.funcargs.get("row") or {}
    case_id = row.get("用例ID") or ""
    if not case_id:
        return
    global _csv_results
    passed = call.excinfo is None
    _csv_results[case_id] = "PASS" if passed else "FAIL"


def pytest_sessionfinish(session, exitstatus):
    """Write 运行结果 back to CSV or Excel by 用例ID."""
    global _case_path, _is_excel, _csv_results
    if not _csv_results or not _case_path or not _case_path.exists():
        return
    if _is_excel:
        try:
            if load_workbook is None:
                return
            wb = load_workbook(_case_path)
            ws = wb.active
            header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
            header = [str(h).strip() if h is not None else "" for h in header]
            col_idx_case_id = None
            col_idx_result = None
            for i, h in enumerate(header):
                if h == "用例ID":
                    col_idx_case_id = i + 1
                if h == "运行结果":
                    col_idx_result = i + 1
            if col_idx_case_id is None:
                wb.close()
                return
            if col_idx_result is None:
                col_idx_result = len(header) + 1
                ws.cell(row=1, column=col_idx_result, value="运行结果")
            for row_idx in range(2, ws.max_row + 1):
                case_id_cell = ws.cell(row=row_idx, column=col_idx_case_id).value
                case_id = "" if case_id_cell is None else str(case_id_cell).strip()
                if case_id in _csv_results:
                    ws.cell(row=row_idx, column=col_idx_result, value=_csv_results[case_id])
            wb.save(_case_path)
        except Exception:
            pass
        return
    # CSV write-back
    csv_path = _case_path
    try:
        with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            fieldnames = [fn.lstrip("\ufeff").strip() for fn in (reader.fieldnames or [])]
            rows = [{k.lstrip("\ufeff").strip(): v for k, v in r.items()} for r in reader]
        if "运行结果" not in fieldnames:
            fieldnames.append("运行结果")
        for row in rows:
            case_id = row.get("用例ID") or ""
            if case_id in _csv_results:
                row["运行结果"] = _csv_results[case_id]
        with open(csv_path, "w", encoding="utf-8", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
    except Exception:
        pass


@pytest.fixture
def row(request):
    """Parametrized row from CSV. Receives dict from parametrize."""
    return request.param if hasattr(request, "param") else {}


def pytest_generate_tests(metafunc):
    """Parametrize test_api_row with CSV/Excel rows when --csv is given."""
    if "row" not in metafunc.fixturenames:
        return
    case_path = metafunc.config.getoption("--csv", default=None)
    if not case_path:
        metafunc.parametrize("row", [{}], indirect=True)
        return
    path = Path(case_path)
    rows = _load_case_rows(path)
    metafunc.parametrize("row", rows if rows else [{}], indirect=True)
