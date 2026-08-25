#!/usr/bin/env python3
"""
export_excel.py（功能测试用例版）
汇总所有批次的测试用例 JSON，导出为标准格式 Excel 文件。
包含三个 Sheet：功能测试用例 / 功能点覆盖维度矩阵 / 覆盖率汇总
"""

import argparse
import json
import glob
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# ── 颜色 ──────────────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="2E75B6") if HAS_OPENPYXL else None
# 按 test_type（内部字段，不导出为列）着色
# test_item 是描述性检查项（如"关键字模糊搜索"），test_type 是覆盖类型（正向/负向/边界/安全）
TYPE_FILLS = {
    "正向": "FFFFFF",
    "负向": "FFF2CC",
    "边界": "DDEEFF",
    "安全": "FFE4CC",
    "权限": "FFE4CC",
} if HAS_OPENPYXL else {}

# ── 标准列定义（25 列，与测试用例管理.xlsx 表头一致，新增关联需求ID、用例版本）─────────
COLUMNS = [
    ("project",         "系统/项目",        15),
    ("module",          "模块",             15),
    ("feature",         "功能",             15),
    ("test_item",       "测试项",           25),   # 功能细化出的具体检查项，描述性文字
    ("title",           "用例标题",         35),
    ("preconditions",   "前置条件",         30),
    ("db_tables",       "涉及到数据库表",   25),
    ("test_data",       "测试数据",         35),
    ("steps",           "测试步骤",         55),
    ("expected_result", "预期结果",         55),
    ("validation_sql",  "验证SQL",          60),
    ("actual_result",   "实际结果",         30),
    ("is_passed",       "是否通过",         10),
    ("estimated_hours", "预计工时(h)",      12),
    ("execute_date",    "执行时间(年月日)", 18),
    ("test_progress",   "测试进度(%)",      12),
    ("case_type",       "用例类型",         12),
    ("priority",        "用例优先级",       12),
    ("is_regression",   "是否回归用例",     13),
    ("is_blocker",      "是否阻塞上线",     13),
    ("requirement_id",  "关联需求ID",       18),   # 禅道需求ID，如 RQ-001
    ("api_url",         "API URL",          30),
    ("bug_count",       "BUG数",             8),
    ("bug_unresolved",  "BUG未解决数",      13),
    ("case_version",    "用例版本",         10),   # 整数版本号：全新功能=1；已有用例基础上新增/修改=原最大版本+1
]
WRAP_COLS = {"steps", "expected_result", "validation_sql", "test_data",
             "preconditions", "actual_result"}

# ── 覆盖维度（功能测试专用）──────────────────────────────────
COVERAGE_DIMS = [
    "正常流程", "数据展示", "边界条件",
    "异常流程", "空数据", "权限控制", "业务规则", "数据一致性"
]

# 关键词 → 维度映射
DIM_KEYWORDS = {
    "正常流程":  ["正向", "正常", "成功", "主流程", "主成功"],
    "数据展示":  ["展示", "显示", "列表", "详情", "查看", "数据准确"],
    "边界条件":  ["边界", "最小", "最大", "空", "最长", "超长", "limit"],
    "异常流程":  ["异常", "错误", "失败", "exception", "超时", "网络"],
    "空数据":    ["空数据", "无数据", "没有数据", "空结果", "空状态", "无结果"],
    "权限控制":  ["权限", "越权", "未登录", "角色", "permission", "auth"],
    "业务规则":  ["业务规则", "约束", "限制", "规则", "判定"],
    "数据一致性":["一致", "验证SQL", "数据库", "DB", "sql"],
}


def infer_coverage_dims(case: dict) -> list[str]:
    """从用例内容推断覆盖的维度"""
    text = " ".join([
        str(case.get("title", "")),
        str(case.get("description", "")),
        str(case.get("case_type", "")),
        str(case.get("design_method", "")),
        str(case.get("expected_result", "")),
        str(case.get("notes", "")),
    ]).lower()

    dims = []
    for dim, kws in DIM_KEYWORDS.items():
        if any(kw.lower() in text for kw in kws):
            dims.append(dim)
    return dims


def load_all_batches(workspace: str) -> list[dict]:
    output_dir = Path(workspace) / "output"
    pattern = str(output_dir / "cases_batch_*.json")
    files = sorted(glob.glob(pattern))
    if not files:
        files = sorted(glob.glob(str(Path(workspace) / "cases_batch_*.json")))

    all_cases = []
    for fp in files:
        with open(fp, encoding="utf-8") as f:
            batch = json.load(f)
        cases = batch if isinstance(batch, list) else batch.get("cases", [])
        all_cases.extend(cases)
        print(f"   📂 加载 {fp}: {len(cases)} 条")

    print(f"✅ 共加载 {len(all_cases)} 条功能测试用例")
    return all_cases


def auto_case_id(cases: list[dict]) -> list[dict]:
    for i, c in enumerate(cases, 1):
        if not c.get("case_id"):
            c["case_id"] = f"TC-FUNC-{i:03d}"
    return cases


def make_thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)


def export_excel(cases: list[dict], output_path: str):
    if not HAS_OPENPYXL:
        raise RuntimeError("请安装 openpyxl: pip install openpyxl")

    wb = openpyxl.Workbook()
    border = make_thin_border()

    # ══════════════════════════════════════════════════════════
    # Sheet 1：功能测试用例
    # ══════════════════════════════════════════════════════════
    ws = wb.active
    ws.title = "功能测试用例"

    hdr_font  = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for ci, (key, name, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(1, ci, name)
        cell.fill = HEADER_FILL
        cell.font = hdr_font
        cell.alignment = hdr_align
        cell.border = border
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    data_font  = Font(name="微软雅黑", size=10)
    left_wrap  = Alignment(horizontal="left", vertical="top", wrap_text=True)
    center_top = Alignment(horizontal="center", vertical="top")

    for ri, case in enumerate(cases, 2):
        # 按覆盖类型（test_type 内部字段，正向/负向/边界/安全）决定行背景色
        # test_item 是描述性检查项文字，test_type 才是类型标签
        ct = case.get("test_type", "正向")
        fill_color = TYPE_FILLS.get(ct, "FFFFFF")
        row_fill = PatternFill("solid", fgColor=fill_color)

        for ci, (key, name, width) in enumerate(COLUMNS, 1):
            val = case.get(key, "") or ""
            cell = ws.cell(ri, ci, str(val))
            cell.fill = row_fill
            cell.font = data_font
            cell.border = border
            if key in WRAP_COLS:
                cell.alignment = left_wrap
            elif key in ("is_passed", "case_type", "priority", "is_regression",
                         "is_blocker", "bug_count", "bug_unresolved",
                         "estimated_hours", "test_progress"):
                cell.alignment = center_top
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top")
        ws.row_dimensions[ri].height = 65

    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(cases)+1}"

    # ══════════════════════════════════════════════════════════
    # Sheet 2：功能点覆盖维度矩阵
    # ══════════════════════════════════════════════════════════
    ws2 = wb.create_sheet("功能点覆盖矩阵")
    _write_coverage_matrix(ws2, cases, border, data_font)

    # ══════════════════════════════════════════════════════════
    # Sheet 3：覆盖率汇总
    # ══════════════════════════════════════════════════════════
    ws3 = wb.create_sheet("覆盖率汇总")
    _write_stats_sheet(ws3, cases, border, data_font)

    wb.save(output_path)
    print(f"✅ Excel 已生成: {output_path}")


def _write_coverage_matrix(ws, cases, border, data_font):
    """功能点 × 覆盖维度矩阵"""
    from openpyxl.styles import PatternFill, Font, Alignment

    # 按「模块/功能点」分组
    feature_groups: dict[str, list[dict]] = {}
    for c in cases:
        key = f"{c.get('module','')}/{c.get('feature','')}"
        feature_groups.setdefault(key, []).append(c)

    title_font = Font(name="微软雅黑", bold=True, size=12)
    hdr_font   = Font(name="微软雅黑", bold=True, color="FFFFFF", size=10)
    check_font = Font(name="微软雅黑", size=11, color="375623")

    ws.cell(1, 1, "功能点覆盖维度矩阵").font = title_font
    ws.cell(2, 1, f"共 {len(feature_groups)} 个功能点  |  {len(cases)} 条用例")

    # 表头行
    header_row = 4
    ws.cell(header_row, 1, "功能点（模块/功能）").font = hdr_font
    ws.cell(header_row, 1).fill = PatternFill("solid", fgColor="2E75B6")
    ws.cell(header_row, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions["A"].width = 30

    for di, dim in enumerate(COVERAGE_DIMS, 2):
        c = ws.cell(header_row, di, dim)
        c.font = hdr_font
        c.fill = PatternFill("solid", fgColor="2E75B6")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(di)].width = 12

    ws.column_dimensions[get_column_letter(len(COVERAGE_DIMS) + 2)].width = 8
    ws.cell(header_row, len(COVERAGE_DIMS) + 2, "用例数").font = hdr_font
    ws.cell(header_row, len(COVERAGE_DIMS) + 2).fill = PatternFill("solid", fgColor="2E75B6")

    covered_fill = PatternFill("solid", fgColor="E2EFDA")
    miss_fill    = PatternFill("solid", fgColor="FCE4D6")

    for ri, (feat_key, feat_cases) in enumerate(feature_groups.items(), header_row + 1):
        all_dims_covered = set()
        for c in feat_cases:
            all_dims_covered.update(infer_coverage_dims(c))

        ws.cell(ri, 1, feat_key).font = data_font
        ws.cell(ri, 1).alignment = Alignment(horizontal="left", vertical="center")

        for di, dim in enumerate(COVERAGE_DIMS, 2):
            if dim in all_dims_covered:
                cell = ws.cell(ri, di, "✓")
                cell.fill = covered_fill
                cell.font = check_font
            else:
                cell = ws.cell(ri, di, "")
                cell.fill = miss_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(ri, len(COVERAGE_DIMS) + 2, len(feat_cases)).font = data_font
        ws.cell(ri, len(COVERAGE_DIMS) + 2).alignment = Alignment(horizontal="center")


def _write_stats_sheet(ws, cases, border, data_font):
    """覆盖率汇总 Sheet"""
    from openpyxl.styles import Font, PatternFill, Alignment

    title_font = Font(name="微软雅黑", bold=True, size=13)
    sub_font   = Font(name="微软雅黑", bold=True, size=11)

    ws.cell(1, 1, "功能测试用例覆盖率汇总").font = title_font
    ws.cell(2, 1, f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12

    total = len(cases)
    item_counts = {}   # 按测试项子类型（正向/负向/边界/安全）统计
    module_counts = {}
    feature_counts = {}
    for c in cases:
        ct = c.get("test_type", "未知")
        item_counts[ct] = item_counts.get(ct, 0) + 1
        mod = c.get("module", "未知")
        module_counts[mod] = module_counts.get(mod, 0) + 1
        feat = f"{c.get('module','')}/{c.get('feature','')}"
        feature_counts[feat] = feature_counts.get(feat, 0) + 1

    row = 4
    ws.cell(row, 1, "一、测试项子类型分布").font = sub_font
    row += 1
    ws.cell(row, 1, "测试项类型"); ws.cell(row, 2, "数量"); ws.cell(row, 3, "占比"); ws.cell(row, 4, "目标")
    row += 1
    targets = {"正向": "≥40%", "负向": "≥30%", "边界": "≥20%", "安全/权限": "≥10%"}
    for ct, cnt in sorted(item_counts.items()):
        ws.cell(row, 1, ct); ws.cell(row, 2, cnt)
        ws.cell(row, 3, f"{cnt/total*100:.1f}%" if total else "0%")
        ws.cell(row, 4, targets.get(ct, "-"))
        row += 1
    ws.cell(row, 1, "合计"); ws.cell(row, 2, total)
    row += 2

    ws.cell(row, 1, "二、模块用例分布").font = sub_font
    row += 1
    ws.cell(row, 1, "模块"); ws.cell(row, 2, "用例数")
    row += 1
    for mod, cnt in sorted(module_counts.items(), key=lambda x: -x[1]):
        ws.cell(row, 1, mod); ws.cell(row, 2, cnt)
        row += 1
    row += 1

    ws.cell(row, 1, "三、功能点用例数分布").font = sub_font
    row += 1
    ws.cell(row, 1, "功能点"); ws.cell(row, 2, "用例数")
    row += 1
    for feat, cnt in sorted(feature_counts.items(), key=lambda x: -x[1]):
        ws.cell(row, 1, feat); ws.cell(row, 2, cnt)
        row += 1


def main():
    parser = argparse.ArgumentParser(description="导出功能测试用例为标准 Excel")
    parser.add_argument("--workspace", required=True, help="工作区路径")
    parser.add_argument("--output",    default=None,  help="输出 Excel 路径")
    parser.add_argument("--type",      default="functional", help="用例类型标识（functional）")
    args = parser.parse_args()

    if not HAS_OPENPYXL:
        print("❌ 请安装 openpyxl: pip install openpyxl")
        return

    cases = load_all_batches(args.workspace)
    if not cases:
        print("⚠️  没有找到 cases_batch_*.json 文件")
        return

    cases = auto_case_id(cases)

    if args.output:
        output_path = args.output
    else:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_dir = Path(args.workspace) / "output"
        out_dir.mkdir(parents=True, exist_ok=True)
        output_path = str(out_dir / f"functional_testcases_{ts}.xlsx")

    export_excel(cases, output_path)

    # 打印覆盖率（按 test_type 内部字段统计）
    total = len(cases)
    item_counts = {}
    for c in cases:
        ct = c.get("test_type", "未知")
        item_counts[ct] = item_counts.get(ct, 0) + 1

    print(f"\n📊 覆盖率汇总（共 {total} 条用例）:")
    targets_map = {"正向": 0.4, "负向": 0.3, "边界": 0.2}
    security_cnt = item_counts.get("安全", 0) + item_counts.get("权限", 0)
    for ct, tgt in targets_map.items():
        cnt = item_counts.get(ct, 0)
        ratio = cnt / total if total else 0
        status = "✅" if ratio >= tgt else "⚠️  未达标"
        print(f"   {status} {ct}: {cnt}/{total} = {ratio*100:.0f}%（目标≥{tgt*100:.0f}%）")
    sec_ratio = security_cnt / total if total else 0
    status = "✅" if sec_ratio >= 0.1 else "⚠️  未达标"
    print(f"   {status} 安全/权限: {security_cnt}/{total} = {sec_ratio*100:.0f}%（目标≥10%）")
    print(f"\n📁 输出文件: {output_path}")


if __name__ == "__main__":
    main()
