#!/usr/bin/env python3
"""
Step 5b: 生成完善后的测试用例文档
基于原测试用例文档，用完善内容覆盖原始单元格，
改动过的单元格高亮（浅黄色），输出干净的完善用例文档（无评审意见列）
"""
import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

SKILL_DIR = Path(__file__).parent.parent

try:
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    import pandas as pd
except ImportError:
    print("缺少依赖: pip install pandas openpyxl", file=sys.stderr)
    sys.exit(1)

# WPS/新版 Excel 兼容补丁
try:
    from openpyxl.worksheet.datavalidation import DataValidation as _DV
    _orig_dv_init = _DV.__init__
    def _patched_dv_init(self, *args, **kwargs):
        kwargs.pop("id", None)
        _orig_dv_init(self, *args, **kwargs)
    _DV.__init__ = _patched_dv_init
except Exception:
    pass


def load_fixed_results(fixed_path: Path) -> tuple[dict[int, dict], str | int | None]:
    """返回 ({row_index: improvements_dict}, sheet名)"""
    with open(fixed_path, encoding="utf-8") as f:
        data = json.load(f)
    improvements_map = {
        r.get("_row_index", i): r.get("_improvements", {})
        for i, r in enumerate(data.get("results", []))
    }
    return improvements_map, data.get("sheet")


def build_col_index_map(ws) -> dict[str, int]:
    """从表头行构建 {列名: 列号(1-based)} 的映射"""
    col_map = {}
    for cell in ws[1]:
        if cell.value:
            col_map[str(cell.value).strip()] = cell.column
    return col_map


def write_fixed_excel(original_path: Path, fixed_path: Path,
                      output_path: Path, sheet: str | int | None = None):
    """基于原Excel，将完善内容写回对应单元格"""
    # 加载完善数据
    with open(fixed_path, encoding="utf-8") as f:
        fixed_data = json.load(f)

    results = fixed_data.get("results", [])

    # 从 parsed_testcases.json 获取 col_map（std → orig）
    parsed_path = fixed_path.parent / "parsed_testcases.json"
    col_map = {}
    if parsed_path.exists():
        with open(parsed_path, encoding="utf-8") as f:
            parsed = json.load(f)
        col_map = parsed.get("col_map", {})  # {std_field: orig_col_name}

    # 加载原 Excel（保留格式）
    wb = load_workbook(original_path)
    if sheet is not None and sheet in wb.sheetnames:
        ws = wb[sheet]
    else:
        ws = wb.active

    # 构建表头 → 列号映射
    header_col_idx = build_col_index_map(ws)

    # 高亮颜色：改动过的单元格用浅黄色
    fill_improved = PatternFill(start_color="FFFACD", end_color="FFFACD", fill_type="solid")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 统计
    total_cells_improved = 0

    for result in results:
        row_idx = result.get("_row_index", -1)
        improvements = result.get("_improvements", {})
        if row_idx < 0 or not improvements:
            continue

        excel_row = row_idx + 2  # pandas 0-based → Excel row (1-based + 1 header)

        for std_field, new_val in improvements.items():
            # 找原始列名
            orig_col = col_map.get(std_field)
            if not orig_col:
                continue
            col_idx = header_col_idx.get(orig_col)
            if not col_idx:
                continue

            cell = ws.cell(row=excel_row, column=col_idx)
            old_val = str(cell.value or "").strip()
            if old_val == str(new_val).strip():
                continue  # 内容未变则不标注

            cell.value = new_val
            cell.fill = fill_improved
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            total_cells_improved += 1

    wb.save(output_path)
    return total_cells_improved


def write_fixed_csv(original_path: Path, fixed_path: Path, output_path: Path):
    """基于原CSV生成完善版本"""
    with open(fixed_path, encoding="utf-8") as f:
        fixed_data = json.load(f)

    results = fixed_data.get("results", [])
    parsed_path = fixed_path.parent / "parsed_testcases.json"
    col_map = {}
    if parsed_path.exists():
        with open(parsed_path, encoding="utf-8") as f:
            parsed = json.load(f)
        col_map = parsed.get("col_map", {})

    for enc in ["utf-8", "gbk", "gb2312"]:
        try:
            df = pd.read_csv(original_path, dtype=str, encoding=enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        df = pd.read_csv(original_path, dtype=str, encoding="utf-8", errors="replace")

    for result in results:
        row_idx = result.get("_row_index", -1)
        improvements = result.get("_improvements", {})
        if row_idx < 0 or not improvements:
            continue
        for std_field, new_val in improvements.items():
            orig_col = col_map.get(std_field)
            if orig_col and orig_col in df.columns:
                df.at[row_idx, orig_col] = new_val

    df.to_csv(output_path, index=False, encoding="utf-8-sig")


def main():
    parser = argparse.ArgumentParser(description="生成完善后的测试用例文档")
    parser.add_argument("--original", required=True, help="原始测试用例文件路径")
    parser.add_argument("--fixed", required=True, help="完善结果JSON路径（fixed_testcases.json）")
    parser.add_argument("--output", default=None, help="输出目录（默认: output/）")
    args = parser.parse_args()

    print(f"\n=== Step 5b: 生成完善后的测试用例文档 ===")

    original_path = Path(args.original)
    if not original_path.is_absolute():
        original_path = SKILL_DIR / original_path

    fixed_path = Path(args.fixed)
    if not fixed_path.is_absolute():
        fixed_path = SKILL_DIR / fixed_path

    output_dir = Path(args.output) if args.output else SKILL_DIR / "output"
    output_dir.mkdir(exist_ok=True)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    stem = original_path.stem
    suffix = original_path.suffix.lower()
    output_suffix = ".xlsx" if suffix in (".xlsx", ".xls") else ".csv"
    output_path = output_dir / f"{stem}_fixed_{ts}{output_suffix}"

    # 加载 sheet 信息
    with open(fixed_path, encoding="utf-8") as f:
        fixed_data = json.load(f)
    sheet = fixed_data.get("sheet")
    improved = fixed_data.get("improved", 0)

    if suffix in (".xlsx", ".xls"):
        total_cells = write_fixed_excel(original_path, fixed_path, output_path, sheet=sheet)
        print(f"  共改动单元格: {total_cells} 个（黄色高亮）")
    else:
        write_fixed_csv(original_path, fixed_path, output_path)

    print(f"  完善用例数: {improved} 条")
    print(f"  完善文档: {output_path}")
    print(f"\n[完成] 完善文档生成成功")


if __name__ == "__main__":
    main()
