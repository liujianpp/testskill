#!/usr/bin/env python3
"""
find_uncovered_features.py
对比功能点列表与已有测试用例，找出没有任何测试用例覆盖的功能点。
输出：features_to_design.json（待设计用例的功能点列表）
"""

import argparse
import json
import re
from pathlib import Path

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# ── 已有用例中可能的功能点相关列名关键词 ──────────────────────────
MODULE_COL_KEYWORDS  = ["所属模块", "module", "模块", "功能模块", "模块名"]
FEATURE_COL_KEYWORDS = ["功能点", "feature", "功能", "子功能", "测试项"]


def load_feature_list(feature_list_path: str) -> list[dict]:
    with open(feature_list_path, encoding="utf-8") as f:
        return json.load(f)


def extract_covered_from_excel(testcase_path: str) -> set[str]:
    """
    从已有测试用例 Excel 中提取已有覆盖的「模块/功能点」组合。
    返回 set，元素格式：  "模块名/功能点名"  或  "功能点名"
    """
    if not testcase_path or not Path(testcase_path).exists():
        return set()

    if not HAS_OPENPYXL:
        print("⚠️  未安装 openpyxl，将跳过已有用例对比（pip install openpyxl）")
        return set()

    wb = openpyxl.load_workbook(testcase_path, read_only=True, data_only=True)
    covered = set()

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            continue

        # ── 在前5行中找表头 ──────────────────────────
        header_row_idx = module_col = feature_col = None
        for i, row in enumerate(rows[:5]):
            row_str = [str(c).strip() if c else "" for c in row]
            for j, cell_str in enumerate(row_str):
                cell_lower = cell_str.lower().replace(" ", "")
                if any(kw.lower().replace(" ", "") in cell_lower for kw in MODULE_COL_KEYWORDS):
                    module_col = j
                    header_row_idx = i
                if any(kw.lower().replace(" ", "") in cell_lower for kw in FEATURE_COL_KEYWORDS):
                    feature_col = j
                    header_row_idx = i
            if module_col is not None or feature_col is not None:
                break

        if header_row_idx is None:
            # 找不到标准列，做全文本模糊匹配（降级）
            for row in rows[1:]:
                for cell in row:
                    if cell and isinstance(cell, str) and len(cell) > 1:
                        covered.add(cell.strip())
            continue

        # ── 读取数据行 ──────────────────────────────────
        for row in rows[header_row_idx + 1:]:
            mod_val = (str(row[module_col]).strip()
                       if module_col is not None and module_col < len(row) and row[module_col]
                       else "")
            feat_val = (str(row[feature_col]).strip()
                        if feature_col is not None and feature_col < len(row) and row[feature_col]
                        else "")

            if mod_val:
                covered.add(mod_val)
            if feat_val:
                covered.add(feat_val)
            if mod_val and feat_val:
                covered.add(f"{mod_val}/{feat_val}")

    wb.close()
    print(f"📊 已有测试用例中提取到 {len(covered)} 个模块/功能点标识")
    return covered


def is_covered(feature: dict, covered: set[str]) -> bool:
    """判断该功能点是否已被覆盖"""
    module  = feature["module"].strip()
    feat    = feature["feature"].strip()
    feat_key = feature["feature_key"].strip()  # "模块/功能点"

    # 精确 key 匹配
    if feat_key in covered:
        return True
    # 功能点名单独匹配
    if feat in covered:
        return True
    # 模糊匹配（允许包含关系）
    for c in covered:
        if feat in c or c in feat:
            return True
    return False


def main():
    parser = argparse.ArgumentParser(description="找出功能清单中没有测试用例的功能点")
    parser.add_argument("--feature_list", required=True, help="feature_list.json 路径")
    parser.add_argument("--testcase",     default=None,  help="已有测试用例 Excel 路径（可选）")
    parser.add_argument("--output",       default=None,  help="输出 JSON 路径")
    parser.add_argument("--target_features", default=None,
                        help="逗号分隔的目标功能点（如 '地址搜索,地址验证'），指定后直接返回这些功能点")
    args = parser.parse_args()

    all_features = load_feature_list(args.feature_list)
    print(f"📋 功能清单共 {len(all_features)} 个功能点")

    # 用户指定目标功能点
    if args.target_features:
        targets = [t.strip() for t in args.target_features.split(",")]
        result = [f for f in all_features
                  if f["feature"] in targets or
                     f["module"] in targets or
                     f["feature_key"] in targets or
                     any(t in f["feature"] or t in f["module"] for t in targets)]
        print(f"🎯 用户指定目标：匹配到 {len(result)} 个功能点")
    else:
        covered = extract_covered_from_excel(args.testcase)

        if covered:
            result = [f for f in all_features if not is_covered(f, covered)]
            covered_count = len(all_features) - len(result)
            print(f"\n🔍 分析结果：")
            print(f"   共 {len(all_features)} 个功能点")
            print(f"   已有用例：{covered_count} 个（跳过）")
            print(f"   待设计：{len(result)} 个")
        else:
            result = all_features
            print(f"\n📋 未提供已有用例，将对所有 {len(all_features)} 个功能点设计用例")

    # 输出待设计功能点列表
    output_path = args.output
    if not output_path:
        output_path = str(Path(args.feature_list).parent / "features_to_design.json")

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    # 打印汇总
    print(f"\n{'='*65}")
    print(f"{'#':<4} {'模块':<15} {'功能点':<20} {'关联表':<20} 是否有描述")
    print(f"{'-'*65}")
    for i, feat in enumerate(result, 1):
        tables = ", ".join(feat["db_tables"])[:19] if feat["db_tables"] else "（无）"
        has_desc = "✅" if feat.get("description") else "⚠️ 缺描述"
        print(f"{i:<4} {feat['module']:<15} {feat['feature']:<20} {tables:<20} {has_desc}")
    print(f"{'='*65}")
    print(f"\n✅ 待设计功能点列表已保存: {output_path}")
    return output_path


if __name__ == "__main__":
    main()
