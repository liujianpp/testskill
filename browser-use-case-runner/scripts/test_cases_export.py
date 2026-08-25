#!/usr/bin/env python3
"""Export test cases from .xlsx (tabular) or Markdown (.md) to JSONL for agent consumption."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

# xlsx column header (Chinese) -> JSON key
XLSX_HEADER_MAP = {
    "用例编号": "id",
    "用例名称": "name",
    "功能模块": "module",
    "模块": "module",
    "优先级": "priority",
    "前置条件": "preconditions",
    "测试步骤": "steps",
    "测试数据": "data",
    "预期结果": "expected",
}

MD_LABELS = {
    "模块": "module",
    "功能模块": "module",
    "用例名称": "name",
    "名称": "name",
    "优先级": "priority",
    "前置条件": "preconditions",
    "测试步骤": "steps",
    "测试数据": "data",
    "预期结果": "expected",
    "用例编号": "id",
}


def _normalize_cell(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def export_xlsx(path: Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise SystemExit("需要 openpyxl：pip install openpyxl 或使用项目 pyproject 依赖") from e

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []

    header_row = [_normalize_cell(c) for c in rows[0]]
    col_index: dict[str, int] = {}
    for i, h in enumerate(header_row):
        if h in XLSX_HEADER_MAP:
            col_index[XLSX_HEADER_MAP[h]] = i

    if "id" not in col_index and "name" not in col_index:
        raise SystemExit(
            f"未识别到必需列：需要表头「用例编号」和/或「用例名称」。实际表头: {header_row}"
        )

    cases: list[dict] = []
    for raw in rows[1:]:
        if not raw or all(_normalize_cell(c) == "" for c in raw):
            continue

        def get(key: str) -> str:
            if key not in col_index:
                return ""
            idx = col_index[key]
            if idx >= len(raw):
                return ""
            return _normalize_cell(raw[idx])

        cid = get("id")
        name = get("name")
        if not cid and not name:
            continue
        if not cid:
            cid = name.split()[0] if name else "ROW"
        if not name:
            name = cid

        cases.append(
            {
                "id": cid,
                "name": name,
                "module": get("module"),
                "priority": get("priority"),
                "preconditions": get("preconditions"),
                "steps": get("steps"),
                "data": get("data"),
                "expected": get("expected"),
                "source": "xlsx",
            }
        )
    return cases


def _parse_md_cases(text: str) -> list[dict]:
    """Parse ## blocks with bullet **label**: value lines."""
    lines = text.splitlines()
    cases: list[dict] = []
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        m = re.match(r"^##\s+(\S+)\s*(.*)$", line)
        if not m:
            i += 1
            continue

        case_id = m.group(1).strip()
        title_rest = m.group(2).strip()
        fields: dict[str, str] = {}
        i += 1
        current_key: str | None = None
        buf: list[str] = []

        def flush():
            nonlocal current_key, buf
            if current_key:
                val = "\n".join(buf).strip() if buf else ""
                if val:
                    fields[current_key] = val
            current_key = None
            buf = []

        while i < len(lines):
            raw = lines[i]
            if re.match(r"^##\s+\S", raw.strip()):
                break
            bullet = re.match(r"^[-*]\s*\*\*([^*]+)\*\*:\s*(.*)$", raw.strip())
            if bullet:
                flush()
                label = bullet.group(1).strip()
                rest = bullet.group(2).strip()
                key = MD_LABELS.get(label)
                if key:
                    current_key = key
                    if rest:
                        buf.append(rest)
                i += 1
                continue
            sub = re.match(r"^\s{2,}[-*]\s+(.*)$", raw)
            if sub and current_key == "steps":
                buf.append(sub.group(1).strip())
                i += 1
                continue
            if current_key and raw.strip() and not raw.strip().startswith("##"):
                buf.append(raw.strip())
                i += 1
                continue
            i += 1

        flush()
        name = fields.pop("name", "") or title_rest or case_id
        cid = fields.pop("id", "") or case_id
        cases.append(
            {
                "id": cid,
                "name": name,
                "module": fields.get("module", ""),
                "priority": fields.get("priority", ""),
                "preconditions": fields.get("preconditions", ""),
                "steps": fields.get("steps", ""),
                "data": fields.get("data", ""),
                "expected": fields.get("expected", ""),
                "source": "md",
            }
        )
    return cases


def export_md(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8")
    return _parse_md_cases(text)


def main() -> None:
    p = argparse.ArgumentParser(description="Export test cases to JSONL")
    sub = p.add_subparsers(dest="cmd", required=True)

    px = sub.add_parser("xlsx", help="From .xlsx workbook")
    px.add_argument("path", type=Path)
    px.add_argument("-o", "--output", type=Path, help="Write JSONL to file (utf-8)")

    pe = sub.add_parser("excel", help="Alias of xlsx (same behavior)")
    pe.add_argument("path", type=Path)
    pe.add_argument("-o", "--output", type=Path, help="Write JSONL to file (utf-8)")

    pm = sub.add_parser("md", help="From .md")
    pm.add_argument("path", type=Path)
    pm.add_argument("-o", "--output", type=Path)

    args = p.parse_args()
    path: Path = args.path
    if not path.is_file():
        raise SystemExit(f"文件不存在: {path}")

    if args.cmd in ("xlsx", "excel"):
        cases = export_xlsx(path)
    else:
        cases = export_md(path)

    out_lines = [json.dumps(c, ensure_ascii=False) for c in cases]
    body = "\n".join(out_lines) + ("\n" if out_lines else "")

    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(body, encoding="utf-8")
    else:
        sys.stdout.write(body)


if __name__ == "__main__":
    main()
