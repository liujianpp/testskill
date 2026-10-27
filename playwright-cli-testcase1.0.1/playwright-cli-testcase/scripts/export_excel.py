#!/usr/bin/env python3
"""Export web test cases JSON to Excel (.xlsx)."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

HEADERS = [
    ("case_id", "用例编号"),
    ("module", "模块"),
    ("title", "用例标题"),
    ("priority", "优先级"),
    ("precondition", "前置条件"),
    ("steps", "测试步骤"),
    ("expected", "预期结果"),
    ("method", "设计方法"),
    ("page", "页面"),
    ("notes", "备注"),
]

COL_WIDTHS = {
    "A": 12,
    "B": 18,
    "C": 36,
    "D": 10,
    "E": 28,
    "F": 40,
    "G": 40,
    "H": 18,
    "I": 20,
    "J": 20,
}


def load_data(path: Path) -> dict:
    with path.open(encoding="utf-8") as f:
        data = json.load(f)
    if not isinstance(data, dict) or "cases" not in data:
        raise ValueError("JSON 需包含 title（可选）与 cases 数组")
    if not isinstance(data["cases"], list):
        raise ValueError("cases 必须为数组")
    return data


def export_excel(data: dict, output: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"

    header_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    for col, (_, label) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = header_font
        cell.alignment = wrap

    for row_idx, case in enumerate(data["cases"], start=2):
        if not isinstance(case, dict):
            raise ValueError(f"cases[{row_idx - 2}] 必须为对象")
        for col, (key, _) in enumerate(HEADERS, start=1):
            value = case.get(key, "")
            if value is None:
                value = ""
            cell = ws.cell(row=row_idx, column=col, value=str(value))
            cell.alignment = wrap

    ws.freeze_panes = "A2"
    for col_letter, width in COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    for i in range(1, len(HEADERS) + 1):
        letter = get_column_letter(i)
        if letter not in COL_WIDTHS:
            ws.column_dimensions[letter].width = 16

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description="将 Web 测试用例 JSON 导出为 Excel")
    parser.add_argument("input_json", type=Path, help="输入 JSON 路径")
    parser.add_argument("output_xlsx", type=Path, help="输出 .xlsx 路径")
    args = parser.parse_args()

    try:
        data = load_data(args.input_json)
        export_excel(data, args.output_xlsx)
    except Exception as exc:  # noqa: BLE001 — CLI surface
        print(f"导出失败: {exc}", file=sys.stderr)
        return 1

    print(f"已导出 Excel: {args.output_xlsx.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
