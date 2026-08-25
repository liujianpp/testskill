#!/usr/bin/env python3
"""Extract plain text from PDF, Word, Markdown, or plain text files."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path


def extract_md_or_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def extract_pdf(path: Path) -> str:
    try:
        import pdfplumber
    except ImportError as exc:
        raise SystemExit(
            "缺少依赖 pdfplumber。请执行: pip install pdfplumber python-docx"
        ) from exc

    parts: list[str] = []
    with pdfplumber.open(path) as pdf:
        if not pdf.pages:
            raise SystemExit(f"PDF 无页面: {path}")
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                parts.append(f"--- 第 {i} 页 ---\n{text}")
    joined = "\n\n".join(parts).strip()
    if not joined:
        raise SystemExit(
            f"未能从 PDF 抽取到文字（可能是扫描件/无文字层/加密）: {path}"
        )
    return joined


def extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise SystemExit(
            "缺少依赖 python-docx。请执行: pip install pdfplumber python-docx"
        ) from exc

    doc = Document(path)
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    joined = "\n".join(parts).strip()
    if not joined:
        raise SystemExit(f"Word 文档无可用文本: {path}")
    return joined


def extract(path: Path) -> str:
    if not path.is_file():
        raise SystemExit(f"文件不存在: {path}")

    suffix = path.suffix.lower()
    if suffix in {".md", ".txt", ".markdown", ".text"}:
        return extract_md_or_txt(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix in {".docx"}:
        return extract_docx(path)
    if suffix == ".doc":
        raise SystemExit("不支持旧版 .doc，请先另存为 .docx")
    raise SystemExit(
        f"不支持的格式: {suffix or '(无扩展名)'}。"
        "支持: .pdf .docx .md .txt"
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="抽取需求文档纯文本")
    parser.add_argument("path", type=Path, help="文档路径")
    args = parser.parse_args()
    text = extract(args.path.expanduser().resolve())
    sys.stdout.write(text)
    if not text.endswith("\n"):
        sys.stdout.write("\n")


if __name__ == "__main__":
    main()
